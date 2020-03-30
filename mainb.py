import random
import math
from utils.tools import truncate
from tp1.TestChiCuadrado import TestChiCuadrado
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference, Series

ACCURACY = 4
quantity_numbers = int(input('Ingrese cantidad de números a generar:'))
subintervals = int(input('Ingrese cantidad de subintervalos:'))
random_numbers = []

# python random
for i in range(quantity_numbers):
    random_numbers.append(truncate(random.random(), ACCURACY))

# Congruales Mixto
# TODO: calcular valores de acuerdo a la cantidad y subintervalos

# example
# random_numbers = [0.15, 0.22, 0.41, 0.65, 0.84,
#                   0.81, 0.62, 0.45, 0.32, 0.07,
#                   0.11, 0.29, 0.58, 0.73, 0.93,
#                   0.97, 0.79, 0.55, 0.35, 0.09,
#                   0.99, 0.51, 0.35, 0.02, 0.19,
#                   0.24, 0.98, 0.10, 0.31, 0.17]
# subintervals = int(math.sqrt(len(random_numbers)))

t = TestChiCuadrado(random_numbers, subintervals, ACCURACY)
t.do_test()


# XLS
wb = Workbook(write_only=True)
ws = wb.create_sheet()
max_row = len(t.labels) + 1

labels_xls = ('Intervalo', 'Frecuencia Observada', 'Frecuencia Esperada')
ws.append(labels_xls)

for i in range(len(t.labels)):
    ws.append((t.labels[i], t.freq_observ[t.key_interval[i]], t.freq_esperadas[t.key_interval[i]]))

chart = BarChart()
chart.type = 'col'
chart.style = 10
chart.height = 20
chart.width = 30
chart.title = 'Prueba Chi Cuadrado'
chart.y_axis.title = 'Frecuencia'
chart.x_axis.title = "Intervalo"

data = Reference(ws, min_col=2, min_row=1, max_row=max_row, max_col=3)
categories = Reference(ws, min_col=1, min_row=2, max_row=max_row)
chart.x_axis.delete = False
chart.y_axis.delete = False
chart.add_data(data, titles_from_data=True)
chart.set_categories(categories)
ws.add_chart(chart, 'E2')

wb.save('frecuencias_por_intervalo.xlsx')
