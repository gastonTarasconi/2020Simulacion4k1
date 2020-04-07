from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference


class ToExcel:

    @staticmethod
    def create_excel(data, filename):
        wb = Workbook(write_only=True)
        ws = wb.create_sheet('ChiCuadrado')
        max_row = len(data.labels) + 1

        labels_xls = ('Intervalo', 'Frecuencia Observada', 'Frecuencia Esperada', 'Estadístico de prueba (C)',
                      'Estadístico de prueba acumulado (CA)')
        ws.append(labels_xls)

        for i in range(len(data.labels)):
            ws.append((data.labels[i][1], data.freq_observ[data.labels[i][0]],
                       data.freq_esperadas[data.labels[i][0]], data.chi_2[i], data.chi_2_ac[i]))

        chart = BarChart()
        chart.type = 'col'
        chart.style = 10
        chart.height = 20
        chart.width = 30
        chart.title = 'Prueba Chi Cuadrado'
        chart.y_axis.title = 'Frecuencia'
        chart.x_axis.title = 'Intervalo'

        chart_data = Reference(ws, min_col=2, min_row=1, max_row=max_row, max_col=3)
        categories = Reference(ws, min_col=1, min_row=2, max_row=max_row)
        chart.x_axis.delete = False
        chart.y_axis.delete = False
        chart.add_data(chart_data, titles_from_data=True)
        chart.set_categories(categories)
        ws.add_chart(chart, 'I2')

        wb.save(f'exports/{filename}.xlsx')
