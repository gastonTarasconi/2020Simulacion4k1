from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.image import Image


class Excel:
    def __init__(self, filename):
        self.filename = filename + '.xlsx'
        self.titles = ('Intervalo', 'Frecuencia Observada', 'Frecuencia Esperada', 'Estadístico de prueba (C)',
                       'Estadístico de prueba acumulado (CA)')

    def do_chi_square(self, _type, labels, observed_frequency, expected_frequency, chi_2, chi_2_ac):
        wb = Workbook(write_only=True)
        ws = wb.create_sheet('ChiCuadrado')
        ws.append(self.titles)

        self._do_chi_square_table(ws, _type, labels, observed_frequency, expected_frequency, chi_2, chi_2_ac)

        self._do_chi_square_graph(ws, 'Prueba Chi Cuadrado', 'Frecuencia', 'Intervalo', labels)

        self._put_chi_square_table(ws)

        wb.save(self.filename)

    def _do_chi_square_table(self, ws, _type, labels, observed_frequencies, expected_frequencies, chi_2, chi_2_ac):
        n_intervals = len(labels)
        for i in range(n_intervals):
            label = str(labels[i][0]) + ' - ' + str(labels[i][1]) if (
                    _type != 3) else str(int(labels[i][0])) + ' - ' + str(int(labels[i][1])) if labels[i][0] != \
                                                                                                labels[i][1] else str(
                int(labels[i][0]))

            ws.append((label,
                       observed_frequencies[i],
                       expected_frequencies[i],
                       chi_2[i],
                       chi_2_ac[i]))
            # ws.append((label, observed_frequency, expected_frequencies[i], chi_2[i], chi_2_ac[i]))

    def _do_chi_square_graph(self, ws, title, y_axis, x_axis, labels):
        chart = BarChart()
        chart.type = 'col'
        chart.style = 10
        chart.height = 20
        chart.width = 30
        chart.title = title
        chart.y_axis.title = y_axis
        chart.x_axis.title = x_axis

        # mas uno por la fila de titulos
        chart_data = Reference(ws, min_col=2, min_row=1, max_row=len(labels) + 1, max_col=3)
        categories = Reference(ws, min_col=1, min_row=2, max_row=len(labels) + 1)
        chart.x_axis.delete = False
        chart.y_axis.delete = False
        chart.add_data(chart_data, titles_from_data=True)
        chart.set_categories(categories)
        ws.add_chart(chart, 'I2')

    def _put_chi_square_table(self, ws):
        img = Image('utils/probabilistic_tables/chi_square_table.png')
        img.anchor = 'I42'
        ws.add_image(img)
